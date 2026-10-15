#!/usr/bin/env python3
"""Checklist 411 — Setembro/14 no padrão do template 433 (logo + rodapé)."""

from __future__ import annotations

import json
import re
import shutil
import struct
import subprocess
import urllib.parse
import urllib.request
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from lxml import etree

from docx_template_scrub import assert_no_template_leak, restore_template_media, scrub_docx

PROJECT = 411
PROJECT_FOLDER_ID = "17E8U2ZehEL4S4hyAWAcMJvIZ2xhm-7ge"
TEMPLATE = Path("/home/eduardo/Downloads/433.docx")
OUT_DIR = Path("/home/eduardo/workspaces/Borderless/411-Azevedo1/docs/checklists-411")
SHOTS = OUT_DIR / "screenshots" / "14"
LEGAL = "https://visao-cce8f.web.app"
EXCLUSAO = f"{LEGAL}/exclusao-conta"
EXPO = (
    "exp://u.expo.dev/be2008b3-c7c7-453f-af2c-2cf82fcfcd56"
    "?runtime-version=1.0.0&channel-name=preview"
)

LOGINS = [
    f"Legal / Privacidade: {LEGAL}",
    f"Exclusão de conta: {EXCLUSAO}",
    f"App (Expo Go): {EXPO}",
    "usuario@gmail.com · consultor@gmail.com · gustavo@gmail.com · consultor2@gmail.com",
    "Senha (todas): 123456",
]

TOPICS = [
    {
        "title": "1. Vencimento informado só com o dia",
        "bullets": [
            "Se o dia ainda não passou no mês atual (incluindo hoje), a conta usa o mês corrente.",
            "Se o dia já passou, o vencimento vai para o próximo mês, inclusive virada de ano.",
            "Dias inexistentes (31/fev, 31/abr) caem no último dia válido, sem data inválida.",
        ],
        "benefits": [
            "Dia 13 informado em 14/09 vira 13/10, e não uma conta vencida em setembro.",
            "O cliente não vê atraso falso por interpretação do mês atual.",
            "A regra fica igual no app, no sync do planning e no lembrete de renda por dia.",
            "Timezone local do aparelho é respeitado na meia-noite do vencimento.",
        ],
    },
    {
        "title": "2. Classificação vencida, vence hoje e a vencer",
        "bullets": [
            "Contas usam a data calculada para overdue / due_today / pending.",
            "O próprio dia do vencimento permanece a vencer (badge “Vence hoje”).",
            "A lista exibe a data completa (DD/MM/AAAA) em vez de só “dia N”.",
        ],
        "benefits": [
            "Status da lista e do planning ficam coerentes.",
            "O cliente distingue o que vence hoje do que já atrasou.",
            "Consultor vê a mesma regra ao abrir as contas do cliente.",
            "Notificações de conta passam a usar a data corrigida.",
        ],
    },
    {
        "title": "3. Push de chat com app em segundo plano ou fechado",
        "bullets": [
            "O app registra o token Expo no perfil e atualiza presença (chat aberto/fechado).",
            "Ao enviar mensagem, o remetente dispara push Expo para quem não está na conversa.",
            "Toque na notificação abre a conversa e, quando possível, a mensagem.",
        ],
        "benefits": [
            "Funciona com tela bloqueada e app fechado, não só com o listener local.",
            "Quem já está na conversa não recebe push duplicado.",
            "Canais Android de mensagens e permissão iOS/APNs entram no bootstrap.",
            "Banner in-app continua só com o app em primeiro plano.",
        ],
    },
    {
        "title": "4. Menu contextual no estilo WhatsApp",
        "bullets": [
            "Segurar a mensagem abre reações, responder e denunciar — a denúncia não dispara sozinha.",
            "Emojis 👍 ❤️ 😂 😮 😢 🙏 ficam visíveis no topo do menu.",
            "Responder mostra prévia da mensagem original acima do campo de texto.",
        ],
        "benefits": [
            "Corrige o comportamento no iPhone em que o long-press ia direto para denúncia.",
            "A interação fica familiar para o usuário final.",
            "A denúncia permanece disponível como ação explícita.",
            "Mensagens próprias também podem receber reação e resposta.",
        ],
    },
    {
        "title": "5. Reações persistidas em tempo real",
        "bullets": [
            "Cada usuário guarda uma reação no documento da mensagem no Firestore.",
            "Tocar de novo no mesmo emoji remove; outro emoji substitui.",
            "Regras do Firestore liberam update só do campo reactions do próprio uid.",
        ],
        "benefits": [
            "Os dois lados da conversa veem a reação na hora.",
            "Não dá para alterar reação de outra pessoa.",
            "A reação fica acoplada visualmente à bolha.",
            "Compatível com iOS e Android via listener já existente.",
        ],
    },
    {
        "title": "6. Menções com @ no bate-papo",
        "bullets": [
            "Ao digitar @, aparece a lista dos participantes daquela conversa, com filtro.",
            "A menção guarda o ID do usuário e destaca o @nome no texto.",
            "Quem foi mencionado recebe push específico e abre a mensagem ao tocar.",
        ],
        "benefits": [
            "Não depende só do nome exibido para identificar a pessoa.",
            "Não é possível mencionar quem não participa do chat.",
            "No chat global, o push vai só para mencionados (evita spam).",
            "A conversa abre posicionada na mensagem quando ela está carregada.",
        ],
    },
    {
        "title": "7. Horário do lembrete “Registre seus gastos”",
        "bullets": [
            "Em Configurações há liga/desliga e escolha de hora/minuto no fuso local.",
            "A preferência fica no perfil do usuário e substitui o agendamento anterior.",
            "O horário fixo das 21:00 deixou de ser a única opção.",
        ],
        "benefits": [
            "Cada cliente escolhe quando quer ser lembrado.",
            "Não gera duas notificações do mesmo lembrete.",
            "Reiniciar o app reaplica o horário salvo.",
            "Desligar o lembrete cancela o agendamento diário.",
        ],
    },
]

HR_XML = (
    '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:v="urn:schemas-microsoft-com:vml" '
    'xmlns:o="urn:schemas-microsoft-com:office:office">'
    "<w:pict>"
    '<v:rect style="width:0;height:1.5pt" o:hr="t" o:hrstd="t" o:hralign="center" '
    'fillcolor="#A0A0A0" stroked="f"/>'
    "</w:pict></w:r>"
)


def get_token() -> str:
    out = subprocess.check_output(
        [
            "gdbus",
            "call",
            "--session",
            "--dest",
            "org.gnome.OnlineAccounts",
            "--object-path",
            "/org/gnome/OnlineAccounts/Accounts/account_1786639014_0",
            "--method",
            "org.gnome.OnlineAccounts.OAuth2Based.GetAccessToken",
        ],
        text=True,
    )
    return re.search(r"'(ya29\.[^']+)'", out).group(1)


def drive_request(method, url, token, data=None, headers=None):
    h = {"Authorization": f"Bearer {token}"}
    if headers:
        h.update(headers)
    req = urllib.request.Request(url, data=data, headers=h, method=method)
    with urllib.request.urlopen(req) as resp:
        raw = resp.read()
        ctype = resp.headers.get("Content-Type", "")
        if "application/json" in ctype:
            return json.loads(raw)
        return raw


def drive_list(parent_id, token):
    files, page = [], None
    while True:
        params = {
            "q": f"'{parent_id}' in parents and trashed=false",
            "fields": "nextPageToken,files(id,name,mimeType)",
            "pageSize": "200",
            "supportsAllDrives": "true",
            "includeItemsFromAllDrives": "true",
        }
        if page:
            params["pageToken"] = page
        data = drive_request(
            "GET",
            "https://www.googleapis.com/drive/v3/files?" + urllib.parse.urlencode(params),
            token,
        )
        files.extend(data.get("files", []))
        page = data.get("nextPageToken")
        if not page:
            break
    return files


def drive_ensure_folder(parent_id, name, token):
    for f in drive_list(parent_id, token):
        if f["name"] == name and f["mimeType"] == "application/vnd.google-apps.folder":
            return f["id"]
    created = drive_request(
        "POST",
        "https://www.googleapis.com/drive/v3/files?supportsAllDrives=true",
        token,
        data=json.dumps(
            {
                "name": name,
                "mimeType": "application/vnd.google-apps.folder",
                "parents": [parent_id],
            }
        ).encode(),
        headers={"Content-Type": "application/json"},
    )
    return created["id"]


def drive_trash(parent_id, names, token):
    for f in drive_list(parent_id, token):
        if f["name"] in names:
            drive_request(
                "PATCH",
                f"https://www.googleapis.com/drive/v3/files/{f['id']}?supportsAllDrives=true",
                token,
                data=json.dumps({"trashed": True}).encode(),
                headers={"Content-Type": "application/json"},
            )


def drive_upload(parent_id, local_path: Path, drive_name: str, token: str):
    boundary = "borderless_upload_boundary"
    meta = json.dumps({"name": drive_name, "parents": [parent_id]})
    file_bytes = local_path.read_bytes()
    mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    body = (
        f"--{boundary}\r\nContent-Type: application/json; charset=UTF-8\r\n\r\n{meta}\r\n"
        f"--{boundary}\r\nContent-Type: {mime}\r\n\r\n"
    ).encode() + file_bytes + f"\r\n--{boundary}--\r\n".encode()
    return drive_request(
        "POST",
        "https://www.googleapis.com/upload/drive/v3/files?uploadType=multipart&supportsAllDrives=true",
        token,
        data=body,
        headers={"Content-Type": f"multipart/related; boundary={boundary}"},
    )


def clear_runs(p):
    for child in list(p._p):
        if child.tag == qn("w:r"):
            p._p.remove(child)


def add_run(p, text: str, *, bold=False, size_pt=15):
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size_pt)
    return run


def add_hr_run(p):
    p._p.append(etree.fromstring(HR_XML))


def add_br(run):
    br = OxmlElement("w:br")
    run._r.append(br)


def set_plain(p, text: str, *, bold=False, size_pt=15):
    clear_runs(p)
    add_run(p, text, bold=bold, size_pt=size_pt)


def rebuild_header(p):
    clear_runs(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = add_run(p, "Checklist", bold=True, size_pt=22)
    add_br(r)
    add_hr_run(p)
    for line in LOGINS:
        r = add_run(p, line, bold=False, size_pt=11)
        add_br(r)
    add_hr_run(p)


def add_page_break_before(p):
    r = OxmlElement("w:r")
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    r.append(br)
    pPr = p._p.find(qn("w:pPr"))
    if pPr is not None:
        pPr.addnext(r)
    else:
        p._p.insert(0, r)


def count_body_drawings(docx_path: Path) -> int:
    with zipfile.ZipFile(docx_path) as z:
        xml = z.read("word/document.xml")
    root = etree.fromstring(xml)
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    return len(root.xpath(".//w:drawing", namespaces=ns))


def assert_center_title_and_images(doc: Document):
    centered = WD_ALIGN_PARAGRAPH.CENTER
    title_ok = False
    for p in doc.paragraphs:
        text = (p.text or "").strip()
        has_drawing = bool(
            p._element.findall(
                ".//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}inline"
            )
        )
        if text == "Checklist" or text.startswith("Checklist\n"):
            assert p.alignment == centered, "título Checklist deve estar CENTRALIZADO"
            title_ok = True
        if text.startswith("Imagem ") or has_drawing:
            assert p.alignment == centered, f"imagem/rótulo deve estar CENTRALIZADO: {text[:40]!r}"
    assert title_ok, "parágrafo do título Checklist não encontrado"


def assert_header_footer_branding(doc: Document):
    h = doc.sections[0].header._element.xml
    f = doc.sections[0].footer._element.xml
    assert h.count("a:blip") >= 1, "header sem logo (blip)"
    assert f.count("a:blip") >= 1, "footer sem rodapé/logo (blip)"


def clean_docx_zip(docx_path: Path):
    scrub_docx(docx_path)
    restore_template_media(docx_path, TEMPLATE)
    scrub_docx(docx_path)
    assert_no_template_leak(docx_path)


def picture_width_cm(img_path: Path) -> Cm:
    """Expo/mobile checklist: prefer 9 cm; crops estreitos usam 4.5 cm."""
    with img_path.open("rb") as fh:
        fh.seek(16)
        w, h = struct.unpack(">II", fh.read(8))
    if w > 0 and h / w > 2.5:
        return Cm(4.5)
    return Cm(9)


def build_checklist() -> Path:
    assert len(TOPICS) >= 7
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    out_docx = OUT_DIR / "Checklist-411-Setembro-14.docx"
    shutil.copy(TEMPLATE, out_docx)
    doc = Document(out_docx)

    # NÃO limpar header/footer — logo e rodapé do template 433
    rebuild_header(doc.paragraphs[0])

    for i, topic in enumerate(TOPICS):
        base = 1 + i * 10
        set_plain(doc.paragraphs[base], topic["title"], bold=True, size_pt=17)
        for j, bullet in enumerate(topic["bullets"][:3]):
            set_plain(doc.paragraphs[base + 1 + j], f"🔹 {bullet}", size_pt=15)
        set_plain(doc.paragraphs[base + 4], "✔️ Benefícios:", size_pt=15)
        for j, benefit in enumerate(topic["benefits"][:4]):
            set_plain(doc.paragraphs[base + 5 + j], benefit, size_pt=15)

    numbered = [p for p in doc.paragraphs if re.match(r"^\d+\.\s+", (p.text or "").strip())]
    assert len(numbered) >= 7, len(numbered)

    labels = json.loads((SHOTS / "labels.json").read_text())
    img_idx = None
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip().startswith("Imagens"):
            img_idx = i
            break
    if img_idx is None:
        raise RuntimeError("Parágrafo Imagens não encontrado no template")

    set_plain(doc.paragraphs[img_idx], "Imagens", bold=True, size_pt=17)
    add_page_break_before(doc.paragraphs[img_idx])

    body = doc.paragraphs[img_idx]._p.getparent()
    after = doc.paragraphs[img_idx]._p.getnext()
    while after is not None:
        nxt = after.getnext()
        if after.tag == qn("w:sectPr"):
            break
        body.remove(after)
        after = nxt

    assert len(labels) == 4
    for n, item in enumerate(labels[:4], start=1):
        img_path = SHOTS / item["file"]
        assert img_path.exists(), img_path
        label = item["label"]
        if label.startswith("Imagem "):
            # labels.json já vem com prefixo
            caption = label
        else:
            caption = f"Imagem {n} — {label}"

        p_label = doc.add_paragraph()
        p_label.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_run(p_label, caption, bold=True, size_pt=12)

        p_pic = doc.add_paragraph()
        p_pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_pic.add_run().add_picture(str(img_path), width=picture_width_cm(img_path))

    doc.save(out_docx)
    clean_docx_zip(out_docx)

    doc2 = Document(out_docx)
    assert_center_title_and_images(doc2)
    assert_header_footer_branding(doc2)
    n = count_body_drawings(out_docx)
    assert n == 4, n

    text = "\n".join(p.text for p in doc2.paragraphs)
    assert "localhost" not in text.lower()
    assert LEGAL in text
    print(f"Generated {out_docx} drawings={n} header/footer OK")
    return out_docx


def ensure_day_folder(month_id: str, token: str) -> str:
    names = ["14 + estimativa", "14"]
    for name in names:
        for f in drive_list(month_id, token):
            if f["name"] == name and f["mimeType"] == "application/vnd.google-apps.folder":
                return f["id"]
    return drive_ensure_folder(month_id, "14", token)


def main():
    out = build_checklist()
    token = get_token()
    year_id = drive_ensure_folder(PROJECT_FOLDER_ID, "2026", token)
    set_id = drive_ensure_folder(year_id, "Setembro", token)
    day_id = ensure_day_folder(set_id, token)
    drive_trash(day_id, [f"Checklist {PROJECT}", f"Checklist {PROJECT}.docx"], token)
    up = drive_upload(day_id, out, f"Checklist {PROJECT}", token)
    print("Uploaded Setembro/", day_id, up.get("id"))
    print(f"https://drive.google.com/file/d/{up['id']}/view")


if __name__ == "__main__":
    main()
