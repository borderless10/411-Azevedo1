#!/usr/bin/env python3
"""Estimativa 411 — somente Robusta 7 dias → Drive (04 + estimativa / Setembro)."""

from __future__ import annotations

import json
import re
import shutil
import subprocess
import urllib.parse
import urllib.request
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

PROJECT = 411
PROJECT_FOLDER_ID = "17E8U2ZehEL4S4hyAWAcMJvIZ2xhm-7ge"
TEMPLATE = Path("/home/eduardo/Downloads/Estimativa_433.docx")
OUT_PATH = Path(
    "/home/eduardo/workspaces/Borderless/411-Azevedo1/docs/Estimativa-411.docx"
)
YEAR = "2026"
MONTH = "Setembro"
DAY_FOLDER = "04 + estimativa"

MVP_TOTAL = 0
ROBUSTA_TOTAL = 7

MODULES = [
    {
        "title": "1- Alertas e Notificações Push",
        "desc": (
            "Levar os lembretes atuais (contas, ciclo e chat) para o celular mesmo com "
            "o app fechado, com preferências claras nas Configurações."
        ),
        "bullets": [
            "Registro de token push (Expo/FCM) por usuário autenticado",
            "Push de mensagem nova no chat consultor ↔ cliente",
            "Push de conta a vencer e lembrete de ciclo/consultoria",
            "Preferências on/off por tipo de alerta nas Configurações",
        ],
        "robusta_bullets": [
            "Canais iOS/Android alinhados aos lembretes locais já existentes",
            "Tratamento de permissão negada e reativação do token",
        ],
        "mvp": 0,
        "robusta": 2,
    },
    {
        "title": "2- Relatórios e Exportação",
        "desc": (
            "Permitir exportar e compartilhar lançamentos e resumo do planejamento "
            "para uso na consultoria e no dia a dia do cliente."
        ),
        "bullets": [
            "Exportação por período (CSV) de gastos e rendas",
            "Resumo do planejamento para compartilhar (Share nativo)",
            "Fluxo do consultor exportando a partir do cliente selecionado",
        ],
        "robusta_bullets": [
            "Filtros de período e categorias no export",
            "Nomes de arquivo e formatação amigáveis para o cliente",
        ],
        "mvp": 0,
        "robusta": 1,
    },
    {
        "title": "3- Painel da Carteira do Consultor",
        "desc": (
            "Evoluir a home do consultor de lista simples para um painel de priorização "
            "com risco e atalhos operacionais."
        ),
        "bullets": [
            "Cards de risco (atraso no Consumo Moderado, contas, desvio do plano)",
            "Ordenação/filtros por prioridade de atendimento",
            "Atalhos para planejamento, chat e visão espectador",
        ],
        "robusta_bullets": [
            "Reaproveitar métricas já calculadas no detalhe do cliente",
            "Empty states e loading claros na carteira",
        ],
        "mvp": 0,
        "robusta": 1,
    },
    {
        "title": "4- Dashboard Planejado × Realizado",
        "desc": (
            "Criar tela analítica dedicada (hoje o “Dashboard” ainda aponta para a Home) "
            "para apoiar a consultoria mensal com comparação ao plano."
        ),
        "bullets": [
            "Tela dedicada com filtros de período",
            "Evolução e desvio por categoria / acompanhados",
            "Visão útil para cliente e consultor (espectador)",
        ],
        "robusta_bullets": [
            "Gráficos alinhados ao padrão visual do app",
            "Indicadores de cumprimento do plano no ciclo atual",
        ],
        "mvp": 0,
        "robusta": 2,
    },
    {
        "title": "5- Estabilidade e UX Crítica",
        "desc": (
            "Fechar pontos parciais da validação (zeros de renda, contas→gasto) e "
            "endurecer os fluxos mais usados no dia a dia."
        ),
        "bullets": [
            "Persistência correta de zeros em renda acompanhada",
            "Consistência contas planejadas → registrar gasto → pagamento",
            "Loading/empty states nos caminhos principais",
            "Checklist de regressão manual dos fluxos críticos",
        ],
        "robusta_bullets": [
            "Correção de IDs legados de gastos esperados",
            "Polimento fino pós feedback de QA",
        ],
        "mvp": 0,
        "robusta": 1,
    },
]

# Template slots (≈6 páras cada); usamos 5 módulos
MODULE_STARTS = [11, 17, 23, 29, 35]
UNUSED_MODULE_RANGES = [
    range(41, 47),  # 6
    range(47, 52),  # 7
    range(52, 57),  # 8
]
RESUMO_HEADER = 57
RESUMO_MVP_TITLE = 58
RESUMO_MVP_BODY = 59
RESUMO_ROBUSTA_TITLE = 60
RESUMO_ROBUSTA_BODY = 61
DISCLAIMER_START = 98


def set_text(p, text: str, bold: bool = False, size: int | None = 15):
    p.text = text
    for run in p.runs:
        run.bold = bold
        if size:
            run.font.size = Pt(size)


def set_label_value(p, label: str, value: str, size: int = 15):
    p.text = ""
    r1 = p.add_run(label)
    r1.bold = True
    r1.font.size = Pt(size)
    r2 = p.add_run(value)
    r2.bold = False
    r2.font.size = Pt(size)


def set_mvp_robusta_lines(p, mvp: int, robusta: int, size: int = 15):
    p.text = ""
    lines = []
    if mvp > 0:
        lines.append(("☑️ MVP: ", f"{mvp} dia(s)"))
    if robusta > 0:
        lines.append(("✅ Robusta: ", f"{robusta} dia(s)"))
    for i, (lab, val) in enumerate(lines):
        if i:
            p.add_run("\n")
        r1 = p.add_run(lab)
        r1.bold = True
        r1.font.size = Pt(size)
        r2 = p.add_run(val)
        r2.bold = False
        r2.font.size = Pt(size)


def add_page_break_before(p):
    r = OxmlElement("w:r")
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    r.append(br)
    pPr = p._p.find(qn("w:pPr"))
    if pPr is not None:
        pPr.addnext(r)
    else:
        p._p.insert(0, r)


def keep_together(p):
    pPr = p._p.get_or_add_pPr()
    for tag in ("w:keepNext", "w:keepLines"):
        el = pPr.find(qn(tag))
        if el is None:
            pPr.append(OxmlElement(tag))


def build_estimativa(out_path: Path):
    modules = [m for m in MODULES if m["mvp"] > 0 or m["robusta"] > 0]
    assert all(m["mvp"] > 0 or m["robusta"] > 0 for m in modules)
    assert sum(m["mvp"] for m in modules) == MVP_TOTAL
    assert sum(m["robusta"] for m in modules) == ROBUSTA_TOTAL
    assert MVP_TOTAL == 0
    assert ROBUSTA_TOTAL == 7

    shutil.copy(TEMPLATE, out_path)
    doc = Document(out_path)

    set_text(
        doc.paragraphs[0],
        "Estimativa em relação ao escopo do projeto",
        bold=True,
        size=17,
    )
    set_text(
        doc.paragraphs[1],
        '(Versão "Robusta" — novas funcionalidades e melhorias do Visão / 411)',
        size=15,
    )
    set_text(doc.paragraphs[4], "1. Visão Geral", bold=True, size=17)
    set_text(
        doc.paragraphs[5],
        "O Visão já opera como companion de consultoria financeira (auth e papéis, "
        "rendas/gastos, Consumo Moderado, contas, cartões, planejamento, metas, "
        "ranking, chat e visão do consultor). Esta estimativa cobre apenas a "
        "evolução Robusta: novas funcionalidades e melhorias de sistema — sem faixa MVP.",
        size=15,
    )
    set_label_value(
        doc.paragraphs[6],
        "Estimativa Robusta: ",
        f"{ROBUSTA_TOTAL} dias úteis — push, exportação, painel do consultor, "
        "dashboard planejado×realizado e estabilidade/UX crítica.",
    )
    set_text(doc.paragraphs[7], "", size=15)
    set_text(doc.paragraphs[8], "", size=15)
    set_text(doc.paragraphs[10], "2. Módulos do Projeto", bold=True, size=17)

    for start, mod in zip(MODULE_STARTS, modules):
        set_text(doc.paragraphs[start], mod["title"], bold=True, size=17)
        set_text(doc.paragraphs[start + 1], mod["desc"], size=15)
        bullets = "• " + "\n• ".join(mod["bullets"])
        set_text(doc.paragraphs[start + 2], bullets, size=15)
        set_text(
            doc.paragraphs[start + 3],
            "✅ Robusta inclui também:\n"
            + "\n".join(f"• {b}" for b in mod["robusta_bullets"]),
            size=15,
        )
        set_mvp_robusta_lines(doc.paragraphs[start + 4], mod["mvp"], mod["robusta"])
        # limpar eventual 6º slot do bloco (quando existir)
        if start + 5 < len(doc.paragraphs) and start + 5 not in MODULE_STARTS:
            # só limpa se for o gap vazio típico entre módulos (ex.: 16, 22…)
            gap = start + 5
            if gap < MODULE_STARTS[-1] or gap < UNUSED_MODULE_RANGES[0].start:
                if not any(gap == s for s in MODULE_STARTS):
                    # gaps 16,22,28,34,40
                    if gap in (16, 22, 28, 34, 40):
                        set_text(doc.paragraphs[gap], "", size=15)

    for rng in UNUSED_MODULE_RANGES:
        for i in rng:
            if i < len(doc.paragraphs):
                set_text(doc.paragraphs[i], "", size=15)

    set_text(doc.paragraphs[RESUMO_HEADER], "3. Resumo das Estimativas", bold=True, size=17)
    set_text(doc.paragraphs[RESUMO_MVP_TITLE], "", size=15)
    set_text(doc.paragraphs[RESUMO_MVP_BODY], "", size=15)
    set_text(
        doc.paragraphs[RESUMO_ROBUSTA_TITLE],
        f"📌 Estimativa Robusta: {ROBUSTA_TOTAL} dias úteis",
        size=15,
    )
    set_text(
        doc.paragraphs[RESUMO_ROBUSTA_BODY],
        "Pacote de evolução do Visão: notificações push, relatórios/exportação, "
        "painel de carteira do consultor, dashboard planejado × realizado e "
        "fechamento de estabilidade/UX nos fluxos críticos — sem faixa MVP.",
        size=15,
    )

    for i in range(RESUMO_ROBUSTA_BODY + 1, DISCLAIMER_START):
        if i < len(doc.paragraphs):
            set_text(doc.paragraphs[i], "", size=15)

    if DISCLAIMER_START < len(doc.paragraphs):
        add_page_break_before(doc.paragraphs[DISCLAIMER_START])
        for i in range(DISCLAIMER_START, min(DISCLAIMER_START + 5, len(doc.paragraphs))):
            keep_together(doc.paragraphs[i])

    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "MVP: 0" not in full_text
    assert "Robusta: 0" not in full_text
    assert "Estimativa MVP:" not in full_text
    assert "📌 Estimativa MVP:" not in full_text
    assert f"Estimativa Robusta: {ROBUSTA_TOTAL} dias" in full_text or (
        f"📌 Estimativa Robusta: {ROBUSTA_TOTAL} dias" in full_text
    )
    # não deixar vazar totais do template 433
    assert "11 dias" not in full_text
    assert "A versão MVP contempla 7 dias" not in full_text

    doc.save(out_path)
    print(f"Saved {out_path}")
    print(
        f"MVP sum={sum(m['mvp'] for m in modules)}, "
        f"Robusta sum={sum(m['robusta'] for m in modules)}"
    )
    for m in modules:
        print(f"  {m['title']}: Robusta {m['robusta']}d")


def get_token() -> str:
    out = subprocess.check_output(
        [
            "gdbus",
            "call",
            "--session",
            "--dest",
            "org.gnome.OnlineAccounts",
            "--object-path",
            "/org/gnome/OnlineAccounts/Accounts/account_1786639014_0",
            "--method",
            "org.gnome.OnlineAccounts.OAuth2Based.GetAccessToken",
        ],
        text=True,
    )
    return re.search(r"'(ya29\.[^']+)'", out).group(1)


def drive_request(method, url, token, data=None, headers=None):
    h = {"Authorization": f"Bearer {token}"}
    if headers:
        h.update(headers)
    req = urllib.request.Request(url, data=data, headers=h, method=method)
    with urllib.request.urlopen(req) as resp:
        body = resp.read()
        return json.loads(body) if body else {}


def drive_list(parent_id, token):
    q = f"'{parent_id}' in parents and trashed=false"
    url = "https://www.googleapis.com/drive/v3/files?" + urllib.parse.urlencode(
        {"q": q, "fields": "files(id,name,mimeType)", "pageSize": 200}
    )
    return drive_request("GET", url, token).get("files", [])


def drive_find(parent_id, name, token):
    for f in drive_list(parent_id, token):
        if f["name"] == name:
            return f["id"]
    return None


def drive_ensure_folder(parent_id, name, token):
    existing = drive_find(parent_id, name, token)
    if existing:
        return existing
    return drive_request(
        "POST",
        "https://www.googleapis.com/drive/v3/files",
        token,
        data=json.dumps(
            {
                "name": name,
                "mimeType": "application/vnd.google-apps.folder",
                "parents": [parent_id],
            }
        ).encode(),
        headers={"Content-Type": "application/json"},
    )["id"]


def drive_rename(file_id, name, token):
    drive_request(
        "PATCH",
        f"https://www.googleapis.com/drive/v3/files/{file_id}?fields=id,name",
        token,
        data=json.dumps({"name": name}).encode(),
        headers={"Content-Type": "application/json"},
    )


def drive_trash_by_id(file_id: str, token: str):
    drive_request(
        "PATCH",
        f"https://www.googleapis.com/drive/v3/files/{file_id}",
        token,
        data=json.dumps({"trashed": True}).encode(),
        headers={"Content-Type": "application/json"},
    )


def drive_ensure_day_folder(month_id: str, token: str) -> str:
    matches = [f for f in drive_list(month_id, token) if f["name"] == DAY_FOLDER]
    if matches:
        keep_id = matches[0]["id"]
        for dup in matches[1:]:
            drive_trash_by_id(dup["id"], token)
            print(f"Trashed pasta duplicada: {DAY_FOLDER} ({dup['id']})")
        return keep_id
    legacy = drive_find(month_id, "04", token)
    if legacy:
        drive_rename(legacy, DAY_FOLDER, token)
        return legacy
    return drive_ensure_folder(month_id, DAY_FOLDER, token)


def drive_trash(parent_id, names, token):
    for f in drive_list(parent_id, token):
        if f["name"] in names:
            drive_trash_by_id(f["id"], token)
            print(f"Trashed {f['name']} ({f['id']})")


def drive_upload(parent_id, local_path: Path, drive_name: str, token: str):
    boundary = "borderless_upload_boundary"
    meta = json.dumps({"name": drive_name, "parents": [parent_id]})
    file_bytes = local_path.read_bytes()
    mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    body = (
        f"--{boundary}\r\nContent-Type: application/json; charset=UTF-8\r\n\r\n{meta}\r\n"
        f"--{boundary}\r\nContent-Type: {mime}\r\n\r\n"
    ).encode() + file_bytes + f"\r\n--{boundary}--\r\n".encode()
    drive_request(
        "POST",
        "https://www.googleapis.com/upload/drive/v3/files?uploadType=multipart",
        token,
        data=body,
        headers={"Content-Type": f"multipart/related; boundary={boundary}"},
    )


def upload(token: str):
    year_id = drive_ensure_folder(PROJECT_FOLDER_ID, YEAR, token)
    month_id = drive_ensure_folder(year_id, MONTH, token)
    day_id = drive_ensure_day_folder(month_id, token)
    drive_trash(
        day_id,
        [f"Estimativa {PROJECT}", f"Estimativa {PROJECT}.docx"],
        token,
    )
    drive_upload(day_id, OUT_PATH, f"Estimativa {PROJECT}", token)
    print(f"Uploaded: {YEAR}/{MONTH}/{DAY_FOLDER}/Estimativa {PROJECT}")
    print(f"Folder id: {day_id}")


if __name__ == "__main__":
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    build_estimativa(OUT_PATH)
    upload(get_token())
