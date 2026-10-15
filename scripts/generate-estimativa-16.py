#!/usr/bin/env python3
"""Estimativa 411 — MVP 5 / Robusta 7 → Drive (16 + estimativa / Setembro)."""

from __future__ import annotations

import json
import re
import shutil
import subprocess
import urllib.parse
import urllib.request
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

PROJECT = 411
PROJECT_FOLDER_ID = "17E8U2ZehEL4S4hyAWAcMJvIZ2xhm-7ge"
TEMPLATE = Path("/home/eduardo/Downloads/Estimativa_433.docx")
OUT_PATH = Path(
    "/home/eduardo/workspaces/Borderless/411-Azevedo1/docs/Estimativa-411-16.docx"
)
YEAR = "2026"
MONTH = "Setembro"
DAY_FOLDER = "16 + estimativa"

MVP_TOTAL = 5
ROBUSTA_TOTAL = 7

MODULES = [
    {
        "title": "1- Contas a pagar: vencida no mês e confirmação de data",
        "desc": (
            "Conta com vencimento só no dia do mês deixa de aparecer como vencida depois "
            "que o dia passa e o sistema empurra para o próximo mês, mesmo sem pagamento. "
            "Reinício de ciclo pode reforçar isso. Ao alterar a data, se o dia já passou "
            "no mês atual, perguntar se a parcela está vencida e se este mês foi pago."
        ),
        "bullets": [
            "Status vencida pelo dia deste mês, sem pular o atraso até marcar paga",
            "Tratar reinício de ciclo sem apagar o atraso em aberto",
            "Ao editar a data, perguntar se já está vencida e se o mês está pago",
        ],
        "robusta_bullets": [
            "Histórico visível de parcelas não pagas entre meses",
        ],
        "mvp": 2,
        "robusta": 2,
    },
    {
        "title": "2- Consultor: scroll Android, renda acompanhada e recomendações",
        "desc": (
            "A lista de clientes no Android não chega ao fim. Renda acompanhada tem tipo "
            "(salário), mas não deixa descrever o caso (cabelo, sobrancelha, rede) como "
            "no Consumo Moderado. No PC/web não dá para excluir recomendação."
        ),
        "bullets": [
            "Scroll da carteira de clientes até o último item no Android",
            "Nome livre na renda acompanhada, além do tipo",
            "Excluir recomendação também no PC/web",
        ],
        "robusta_bullets": [
            "Ajustes visuais nessas telas para deixar o fluxo mais óbvio",
        ],
        "mvp": 1,
        "robusta": 1,
    },
    {
        "title": "3- Avisos do ciclo com antecedência",
        "desc": (
            "A notificação para agendar consultoria só aparece quando o ciclo já acabou. "
            "Falta avisar antes, para o cliente marcar com tempo."
        ),
        "bullets": [
            "Aviso 7, 5 e 3 dias antes do fim do ciclo",
            "Aviso no dia anterior e no dia em que termina",
            "Sem repetir o mesmo aviso no mesmo dia",
        ],
        "robusta_bullets": [
            "Texto e horário dos avisos mais claros na home",
        ],
        "mvp": 1,
        "robusta": 1,
    },
    {
        "title": "4- Publicação App Store e Google Play",
        "desc": (
            "As correções precisam chegar no app das lojas. Inclui build de produção, "
            "envio do AAB e do IPA e publicação da versão."
        ),
        "bullets": [
            "Build de produção Android (AAB) e iOS",
            "Upload e lançamento no Google Play e na App Store",
            "Conferência da versão publicada nos dois canais",
        ],
        "robusta_bullets": [
            "Passada extra de QA nas lojas depois das melhorias visuais",
        ],
        "mvp": 1,
        "robusta": 1,
    },
    {
        "title": "5- Intuitividade visual e índices de performance",
        "desc": (
            "Na versão robusta, além do MVP, entram ajustes visuais para o app ficar "
            "mais intuitivo e a criação de índices no Firestore para consultas mais rápidas."
        ),
        "bullets": [
            "Revisão visual das telas de contas, renda e clientes",
            "Índices Firestore nas consultas mais usadas",
            "Medir o tempo de abertura das listas principais",
        ],
        "robusta_bullets": [
            "Polimento de rótulos, hierarquia e estados vazios",
            "Índices extras nas listas do consultor e do planejamento",
        ],
        "mvp": 0,
        "robusta": 2,
    },
]


MODULE_STARTS = [11, 17, 23, 29, 35]
UNUSED_MODULE_RANGES = [
    range(41, 57),
]
RESUMO_HEADER = 57
RESUMO_MVP_TITLE = 58
RESUMO_MVP_BODY = 59
RESUMO_ROBUSTA_TITLE = 60
RESUMO_ROBUSTA_BODY = 61
DISCLAIMER_START = 98


def set_text(p, text: str, bold: bool = False, size: int | None = 15):
    p.text = text
    for run in p.runs:
        run.bold = bold
        if size:
            run.font.size = Pt(size)


def set_label_value(p, label: str, value: str, size: int = 15):
    p.text = ""
    r1 = p.add_run(label)
    r1.bold = True
    r1.font.size = Pt(size)
    r2 = p.add_run(value)
    r2.bold = False
    r2.font.size = Pt(size)


def set_mvp_robusta_lines(p, mvp: int, robusta: int, size: int = 15):
    p.text = ""
    lines = []
    if mvp > 0:
        lines.append(("☑️ MVP: ", f"{mvp} dia(s)"))
    if robusta > 0:
        lines.append(("✅ Robusta: ", f"{robusta} dia(s)"))
    for i, (lab, val) in enumerate(lines):
        if i:
            p.add_run("\n")
        r1 = p.add_run(lab)
        r1.bold = True
        r1.font.size = Pt(size)
        r2 = p.add_run(val)
        r2.bold = False
        r2.font.size = Pt(size)


def add_page_break_before(p):
    r = OxmlElement("w:r")
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    r.append(br)
    pPr = p._p.find(qn("w:pPr"))
    if pPr is not None:
        pPr.addnext(r)
    else:
        p._p.insert(0, r)


def keep_together(p):
    pPr = p._p.get_or_add_pPr()
    for tag in ("w:keepNext", "w:keepLines"):
        el = pPr.find(qn(tag))
        if el is None:
            pPr.append(OxmlElement(tag))


def build_estimativa(out_path: Path):
    modules = [m for m in MODULES if m["mvp"] > 0 or m["robusta"] > 0]
    assert all(m["mvp"] > 0 or m["robusta"] > 0 for m in modules)
    assert sum(m["mvp"] for m in modules) == MVP_TOTAL
    assert sum(m["robusta"] for m in modules) == ROBUSTA_TOTAL
    assert MVP_TOTAL == 5
    assert ROBUSTA_TOTAL == 7

    shutil.copy(TEMPLATE, out_path)
    doc = Document(out_path)

    set_text(
        doc.paragraphs[0],
        "Estimativa em relação ao escopo do projeto",
        bold=True,
        size=17,
    )
    set_text(
        doc.paragraphs[1],
        '(Versões "MVP" e "Robusta" — contas vencidas, consultor, ciclo e publicação nas lojas)',
        size=15,
    )
    set_text(doc.paragraphs[4], "1. Visão Geral", bold=True, size=17)
    set_text(
        doc.paragraphs[5],
        "O Visão já opera planejamento, contas, chat e consultoria. Esta estimativa cobre "
        "o que ainda falta: conta do mês ficar vencida até pagar, confirmação ao alterar "
        "data, scroll de clientes no Android, nome livre na renda acompanhada, excluir "
        "recomendação no PC, avisos do ciclo com antecedência e publicação na App Store "
        "e no Google Play.",
        size=15,
    )
    set_label_value(
        doc.paragraphs[6],
        "Estimativa MVP: ",
        f"{MVP_TOTAL} dias úteis — contas vencidas, ajustes do consultor, avisos do ciclo "
        "e publicação nas lojas.",
    )
    set_label_value(
        doc.paragraphs[7],
        "Estimativa Robusta: ",
        f"{ROBUSTA_TOTAL} dias úteis — o MVP mais alterações visuais para melhorar a "
        "intuitividade do app e criação de índices para deixar as consultas mais rápidas.",
    )
    set_text(doc.paragraphs[8], "", size=15)
    set_text(doc.paragraphs[10], "2. Módulos do Projeto", bold=True, size=17)

    used_starts = set()
    for start, mod in zip(MODULE_STARTS, modules):
        used_starts.add(start)
        set_text(doc.paragraphs[start], mod["title"], bold=True, size=17)
        set_text(doc.paragraphs[start + 1], mod["desc"], size=15)
        set_text(doc.paragraphs[start + 2], "• " + "\n• ".join(mod["bullets"]), size=15)
        extra = mod.get("robusta_bullets") or []
        if extra and mod["robusta"] > mod["mvp"]:
            set_text(
                doc.paragraphs[start + 3],
                "✅ Robusta inclui também:\n" + "\n".join(f"• {b}" for b in extra),
                size=15,
            )
        elif extra and mod["mvp"] == 0:
            set_text(
                doc.paragraphs[start + 3],
                "✅ Robusta inclui também:\n" + "\n".join(f"• {b}" for b in extra),
                size=15,
            )
        else:
            set_text(
                doc.paragraphs[start + 3],
                "✅ Robusta inclui também:\n" + "\n".join(f"• {b}" for b in extra)
                if extra
                else "",
                size=15,
            )
        set_mvp_robusta_lines(doc.paragraphs[start + 4], mod["mvp"], mod["robusta"])

    for i in range(11, RESUMO_HEADER):
        if i in used_starts or any(i in range(s, s + 5) for s in MODULE_STARTS[: len(modules)]):
            continue
        set_text(doc.paragraphs[i], "", size=15)

    set_text(doc.paragraphs[RESUMO_HEADER], "3. Resumo das Estimativas", bold=True, size=17)
    set_text(
        doc.paragraphs[RESUMO_MVP_TITLE],
        f"📌 Estimativa MVP: {MVP_TOTAL} dias úteis",
        size=15,
    )
    set_text(
        doc.paragraphs[RESUMO_MVP_BODY],
        "Conta vencida até pagar, confirmação ao alterar data, scroll de clientes, nome "
        "na renda acompanhada, exclusão de recomendação no PC, avisos 7/5/3/1/0 do ciclo "
        "e publicação na App Store e no Google Play.",
        size=15,
    )
    set_text(
        doc.paragraphs[RESUMO_ROBUSTA_TITLE],
        f"📌 Estimativa Robusta: {ROBUSTA_TOTAL} dias úteis",
        size=15,
    )
    set_text(
        doc.paragraphs[RESUMO_ROBUSTA_BODY],
        "O pacote MVP, com alterações visuais para melhorar a intuitividade do app e "
        "criação de índices para otimizar consultas e deixar o app mais rápido.",
        size=15,
    )

    for i in range(RESUMO_ROBUSTA_BODY + 1, DISCLAIMER_START):
        if i < len(doc.paragraphs):
            set_text(doc.paragraphs[i], "", size=15)

    if DISCLAIMER_START < len(doc.paragraphs):
        add_page_break_before(doc.paragraphs[DISCLAIMER_START])
        for i in range(DISCLAIMER_START, min(DISCLAIMER_START + 5, len(doc.paragraphs))):
            keep_together(doc.paragraphs[i])

    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "MVP: 0" not in full_text
    assert "Robusta: 0" not in full_text
    assert f"{MVP_TOTAL} dias" in full_text
    assert f"{ROBUSTA_TOTAL} dias" in full_text
    assert "11 dias" not in full_text
    assert "A versão MVP contempla 7 dias" not in full_text
    assert "App Store" in full_text
    assert "Google Play" in full_text
    assert "índices" in full_text.lower() or "indices" in full_text.lower()
    assert "intuitiv" in full_text.lower()

    doc.save(out_path)
    print(f"Saved {out_path}")
    print(
        f"MVP sum={sum(m['mvp'] for m in modules)}, "
        f"Robusta sum={sum(m['robusta'] for m in modules)}"
    )
    for m in modules:
        print(f"  {m['title']}: MVP {m['mvp']}d / Robusta {m['robusta']}d")


def get_token() -> str:
    out = subprocess.check_output(
        [
            "gdbus",
            "call",
            "--session",
            "--dest",
            "org.gnome.OnlineAccounts",
            "--object-path",
            "/org/gnome/OnlineAccounts/Accounts/account_1786639014_0",
            "--method",
            "org.gnome.OnlineAccounts.OAuth2Based.GetAccessToken",
        ],
        text=True,
    )
    return re.search(r"'(ya29\.[^']+)'", out).group(1)


def drive_request(method, url, token, data=None, headers=None):
    h = {"Authorization": f"Bearer {token}"}
    if headers:
        h.update(headers)
    req = urllib.request.Request(url, data=data, headers=h, method=method)
    with urllib.request.urlopen(req) as resp:
        body = resp.read()
        return json.loads(body) if body else {}


def drive_list(parent_id, token):
    q = f"'{parent_id}' in parents and trashed=false"
    url = "https://www.googleapis.com/drive/v3/files?" + urllib.parse.urlencode(
        {
            "q": q,
            "fields": "files(id,name,mimeType)",
            "pageSize": 200,
            "supportsAllDrives": "true",
            "includeItemsFromAllDrives": "true",
        }
    )
    return drive_request("GET", url, token).get("files", [])


def drive_find(parent_id, name, token):
    for f in drive_list(parent_id, token):
        if f["name"] == name:
            return f["id"]
    return None


def drive_ensure_folder(parent_id, name, token):
    existing = drive_find(parent_id, name, token)
    if existing:
        return existing
    return drive_request(
        "POST",
        "https://www.googleapis.com/drive/v3/files?supportsAllDrives=true",
        token,
        data=json.dumps(
            {
                "name": name,
                "mimeType": "application/vnd.google-apps.folder",
                "parents": [parent_id],
            }
        ).encode(),
        headers={"Content-Type": "application/json"},
    )["id"]


def drive_rename(file_id, name, token):
    drive_request(
        "PATCH",
        f"https://www.googleapis.com/drive/v3/files/{file_id}?supportsAllDrives=true&fields=id,name",
        token,
        data=json.dumps({"name": name}).encode(),
        headers={"Content-Type": "application/json"},
    )


def drive_trash_by_id(file_id: str, token: str):
    drive_request(
        "PATCH",
        f"https://www.googleapis.com/drive/v3/files/{file_id}?supportsAllDrives=true",
        token,
        data=json.dumps({"trashed": True}).encode(),
        headers={"Content-Type": "application/json"},
    )


def drive_ensure_day_folder(month_id: str, token: str) -> str:
    matches = [f for f in drive_list(month_id, token) if f["name"] == DAY_FOLDER]
    if matches:
        return matches[0]["id"]
    legacy = drive_find(month_id, "16", token)
    if legacy:
        drive_rename(legacy, DAY_FOLDER, token)
        return legacy
    return drive_ensure_folder(month_id, DAY_FOLDER, token)


def drive_trash(parent_id, names, token):
    for f in drive_list(parent_id, token):
        if f["name"] in names:
            drive_trash_by_id(f["id"], token)
            print(f"Trashed {f['name']} ({f['id']})")


def drive_upload(parent_id, local_path: Path, drive_name: str, token: str):
    boundary = "borderless_upload_boundary"
    meta = json.dumps({"name": drive_name, "parents": [parent_id]})
    file_bytes = local_path.read_bytes()
    mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    body = (
        f"--{boundary}\r\nContent-Type: application/json; charset=UTF-8\r\n\r\n{meta}\r\n"
        f"--{boundary}\r\nContent-Type: {mime}\r\n\r\n"
    ).encode() + file_bytes + f"\r\n--{boundary}--\r\n".encode()
    return drive_request(
        "POST",
        "https://www.googleapis.com/upload/drive/v3/files?uploadType=multipart&supportsAllDrives=true",
        token,
        data=body,
        headers={"Content-Type": f"multipart/related; boundary={boundary}"},
    )


def upload(token: str):
    year_id = drive_ensure_folder(PROJECT_FOLDER_ID, YEAR, token)
    month_id = drive_ensure_folder(year_id, MONTH, token)
    day_id = drive_ensure_day_folder(month_id, token)
    drive_trash(
        day_id,
        [f"Estimativa {PROJECT}", f"Estimativa {PROJECT}.docx"],
        token,
    )
    up = drive_upload(day_id, OUT_PATH, f"Estimativa {PROJECT}", token)
    print(f"Uploaded: {YEAR}/{MONTH}/{DAY_FOLDER}/Estimativa {PROJECT}")
    print(f"Folder id: {day_id}")
    if isinstance(up, dict) and up.get("id"):
        print(f"https://drive.google.com/file/d/{up['id']}/view")
        print(f"https://drive.google.com/drive/folders/{day_id}")


if __name__ == "__main__":
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    build_estimativa(OUT_PATH)
    upload(get_token())
