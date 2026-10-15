#!/usr/bin/env python3
"""Checklist 411 — Setembro/01 no padrão do template 433 (logo + rodapé)."""

from __future__ import annotations

import json
import re
import shutil
import struct
import subprocess
import urllib.parse
import urllib.request
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from lxml import etree

from docx_template_scrub import assert_no_template_leak, restore_template_media, scrub_docx

PROJECT = 411
PROJECT_FOLDER_ID = "17E8U2ZehEL4S4hyAWAcMJvIZ2xhm-7ge"
TEMPLATE = Path("/home/eduardo/Downloads/433.docx")
OUT_DIR = Path("/home/eduardo/workspaces/Borderless/411-Azevedo1/docs/checklists-411")
SHOTS = OUT_DIR / "screenshots" / "01"
LEGAL = "https://visao-cce8f.web.app"
EXCLUSAO = f"{LEGAL}/exclusao-conta"
EXPO = (
    "exp://u.expo.dev/be2008b3-c7c7-453f-af2c-2cf82fcfcd56"
    "?runtime-version=1.0.0&channel-name=preview"
)

LOGINS = [
    f"Legal / Privacidade: {LEGAL}",
    f"Exclusão de conta: {EXCLUSAO}",
    f"App (Expo Go): {EXPO}",
    "usuario@gmail.com · consultor@gmail.com · gustavo@gmail.com · consultor2@gmail.com",
    "Senha (todas): 123456",
]

TOPICS = [
    {
        "title": "1. Página pública de exclusão de conta e dados",
        "bullets": [
            "Criamos legal/exclusao-conta.html com fluxo LGPD para solicitar exclusão.",
            "Publicada no Hosting legal (visao-cce8f.web.app/exclusao-conta).",
            "Inclui canais (e-mail), confirmação de identidade e prazos esperados.",
        ],
        "benefits": [
            "Atende exigência de lojas (Google/Apple) para exclusão de conta.",
            "Usuário encontra o processo fora do app, em URL pública.",
            "Documentação alinhada à marca Visão / Azevedo.",
            "Reduz risco de rejeição na revisão das stores.",
        ],
    },
    {
        "title": "2. Link de exclusão na Política de Privacidade",
        "bullets": [
            "Atualizamos legal/index.html com item de navegação Exclusão de conta.",
            "Incluímos referência textual apontando para a página de exclusão.",
            "Mantém Privacidade e Termos no mesmo hub legal.",
        ],
        "benefits": [
            "Privacidade deixa de ficar isolada do pedido de exclusão.",
            "Auditoria ANPD/loja encontra o caminho em um clique.",
            "Consistência entre documentos legais.",
            "Menos suporte manual para orientar o cliente.",
        ],
    },
    {
        "title": "3. Link de exclusão nos Termos de Uso",
        "bullets": [
            "Atualizamos legal/termos.html com o mesmo atalho de exclusão.",
            "Navegação entre Privacidade, Termos e Exclusão unificada.",
            "Mesmo padrão visual das demais páginas legais.",
        ],
        "benefits": [
            "Termos não omitem o direito de exclusão.",
            "Usuário chega à instrução de qualquer documento legal.",
            "Experiência previsível no site legal.",
            "Fecha o conjunto mínimo de páginas para compliance.",
        ],
    },
    {
        "title": "4. Atalho no app (Configurações → Termos e Privacidade)",
        "bullets": [
            "No Expo, Configurações abre Privacidade, Termos e Exclusão de conta.",
            "Usa Linking.openURL para as páginas legais em produção.",
            "Subtitle atualizado: privacidade, termos e exclusão de conta.",
        ],
        "benefits": [
            "Cliente acessa compliance sem sair do fluxo do app.",
            "Fecha o gap entre páginas HTML e a navegação Expo.",
            "Evidência visual no checklist com prints do app.",
            "Alinha entrega mobile ao padrão Borderless.",
        ],
    },
    {
        "title": "5. Feature graphic 1024×500 para a Play Store",
        "bullets": [
            "Adicionamos assets/feature-graphic-1024x500.png no padrão Google Play.",
            "Arte alinhada à identidade Visão.",
            "Arquivo pronto para upload no console da loja.",
        ],
        "benefits": [
            "Cumpre requisito visual da ficha do app Android.",
            "Apresentação consistente na listagem da loja.",
            "Acelera o fechamento do pacote de assets.",
            "Evita rejeição por material gráfico incompleto.",
        ],
    },
    {
        "title": "6. Feature graphic Visão AI e ícone 512",
        "bullets": [
            "Incluímos assets/feature-graphic-visao-ai.png para peças de vitrine.",
            "Incluímos assets/icon-512.png para ícone de alta resolução.",
            "Conjunto complementa o material já usado no app.",
        ],
        "benefits": [
            "Material extra para store listing e marketing.",
            "Ícone nítido em dispositivos e fichas.",
            "Padroniza entregáveis de marca do 411.",
            "Reduz retrabalho na publicação.",
        ],
    },
    {
        "title": "7. Checklist do gap pós-emergencial (26/08)",
        "bullets": [
            "Itens acima não estavam no Checklist 26 + emergencial.",
            "Evidências com 4 imagens capturadas de dentro do app Expo.",
            "Pasta Drive Setembro/01 criada para registrar a entrega.",
        ],
        "benefits": [
            "Fecha o relato faltante perante a fila.",
            "Cliente vê o que foi feito depois do emergencial.",
            "Auditoria Drive fica rastreável por dia.",
            "Fotos no padrão mobile Expo (não páginas HTML soltas).",
        ],
    },
]

HR_XML = (
    '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:v="urn:schemas-microsoft-com:vml" '
    'xmlns:o="urn:schemas-microsoft-com:office:office">'
    "<w:pict>"
    '<v:rect style="width:0;height:1.5pt" o:hr="t" o:hrstd="t" o:hralign="center" '
    'fillcolor="#A0A0A0" stroked="f"/>'
    "</w:pict></w:r>"
)


def get_token() -> str:
    out = subprocess.check_output(
        [
            "gdbus",
            "call",
            "--session",
            "--dest",
            "org.gnome.OnlineAccounts",
            "--object-path",
            "/org/gnome/OnlineAccounts/Accounts/account_1786639014_0",
            "--method",
            "org.gnome.OnlineAccounts.OAuth2Based.GetAccessToken",
        ],
        text=True,
    )
    return re.search(r"'(ya29\.[^']+)'", out).group(1)


def drive_request(method, url, token, data=None, headers=None):
    h = {"Authorization": f"Bearer {token}"}
    if headers:
        h.update(headers)
    req = urllib.request.Request(url, data=data, headers=h, method=method)
    with urllib.request.urlopen(req) as resp:
        raw = resp.read()
        ctype = resp.headers.get("Content-Type", "")
        if "application/json" in ctype:
            return json.loads(raw)
        return raw


def drive_list(parent_id, token):
    files, page = [], None
    while True:
        params = {
            "q": f"'{parent_id}' in parents and trashed=false",
            "fields": "nextPageToken,files(id,name,mimeType)",
            "pageSize": "200",
            "supportsAllDrives": "true",
            "includeItemsFromAllDrives": "true",
        }
        if page:
            params["pageToken"] = page
        data = drive_request(
            "GET",
            "https://www.googleapis.com/drive/v3/files?" + urllib.parse.urlencode(params),
            token,
        )
        files.extend(data.get("files", []))
        page = data.get("nextPageToken")
        if not page:
            break
    return files


def drive_ensure_folder(parent_id, name, token):
    for f in drive_list(parent_id, token):
        if f["name"] == name and f["mimeType"] == "application/vnd.google-apps.folder":
            return f["id"]
    created = drive_request(
        "POST",
        "https://www.googleapis.com/drive/v3/files?supportsAllDrives=true",
        token,
        data=json.dumps(
            {
                "name": name,
                "mimeType": "application/vnd.google-apps.folder",
                "parents": [parent_id],
            }
        ).encode(),
        headers={"Content-Type": "application/json"},
    )
    return created["id"]


def drive_trash(parent_id, names, token):
    for f in drive_list(parent_id, token):
        if f["name"] in names:
            drive_request(
                "PATCH",
                f"https://www.googleapis.com/drive/v3/files/{f['id']}?supportsAllDrives=true",
                token,
                data=json.dumps({"trashed": True}).encode(),
                headers={"Content-Type": "application/json"},
            )


def drive_upload(parent_id, local_path: Path, drive_name: str, token: str):
    boundary = "borderless_upload_boundary"
    meta = json.dumps({"name": drive_name, "parents": [parent_id]})
    file_bytes = local_path.read_bytes()
    mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    body = (
        f"--{boundary}\r\nContent-Type: application/json; charset=UTF-8\r\n\r\n{meta}\r\n"
        f"--{boundary}\r\nContent-Type: {mime}\r\n\r\n"
    ).encode() + file_bytes + f"\r\n--{boundary}--\r\n".encode()
    return drive_request(
        "POST",
        "https://www.googleapis.com/upload/drive/v3/files?uploadType=multipart&supportsAllDrives=true",
        token,
        data=body,
        headers={"Content-Type": f"multipart/related; boundary={boundary}"},
    )


def clear_runs(p):
    for child in list(p._p):
        if child.tag == qn("w:r"):
            p._p.remove(child)


def add_run(p, text: str, *, bold=False, size_pt=15):
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size_pt)
    return run


def add_hr_run(p):
    p._p.append(etree.fromstring(HR_XML))


def add_br(run):
    br = OxmlElement("w:br")
    run._r.append(br)


def set_plain(p, text: str, *, bold=False, size_pt=15):
    clear_runs(p)
    add_run(p, text, bold=bold, size_pt=size_pt)


def rebuild_header(p):
    clear_runs(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = add_run(p, "Checklist", bold=True, size_pt=22)
    add_br(r)
    add_hr_run(p)
    for line in LOGINS:
        r = add_run(p, line, bold=False, size_pt=11)
        add_br(r)
    add_hr_run(p)


def add_page_break_before(p):
    r = OxmlElement("w:r")
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    r.append(br)
    pPr = p._p.find(qn("w:pPr"))
    if pPr is not None:
        pPr.addnext(r)
    else:
        p._p.insert(0, r)


def count_body_drawings(docx_path: Path) -> int:
    with zipfile.ZipFile(docx_path) as z:
        xml = z.read("word/document.xml")
    root = etree.fromstring(xml)
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    return len(root.xpath(".//w:drawing", namespaces=ns))


def assert_center_title_and_images(doc: Document):
    centered = WD_ALIGN_PARAGRAPH.CENTER
    title_ok = False
    for p in doc.paragraphs:
        text = (p.text or "").strip()
        has_drawing = bool(
            p._element.findall(
                ".//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}inline"
            )
        )
        if text == "Checklist" or text.startswith("Checklist\n"):
            assert p.alignment == centered, "título Checklist deve estar CENTRALIZADO"
            title_ok = True
        if text.startswith("Imagem ") or has_drawing:
            assert p.alignment == centered, f"imagem/rótulo deve estar CENTRALIZADO: {text[:40]!r}"
    assert title_ok, "parágrafo do título Checklist não encontrado"


def assert_header_footer_branding(doc: Document):
    h = doc.sections[0].header._element.xml
    f = doc.sections[0].footer._element.xml
    assert h.count("a:blip") >= 1, "header sem logo (blip)"
    assert f.count("a:blip") >= 1, "footer sem rodapé/logo (blip)"


def clean_docx_zip(docx_path: Path):
    scrub_docx(docx_path)
    restore_template_media(docx_path, TEMPLATE)
    scrub_docx(docx_path)
    assert_no_template_leak(docx_path)


def picture_width_cm(img_path: Path) -> Cm:
    """Expo/mobile checklist: prefer 9 cm; crops estreitos usam 4.5 cm."""
    with img_path.open("rb") as fh:
        fh.seek(16)
        w, h = struct.unpack(">II", fh.read(8))
    if w > 0 and h / w > 2.5:
        return Cm(4.5)
    return Cm(9)


def build_checklist() -> Path:
    assert len(TOPICS) >= 7
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    out_docx = OUT_DIR / "Checklist-411-Setembro-01.docx"
    shutil.copy(TEMPLATE, out_docx)
    doc = Document(out_docx)

    # NÃO limpar header/footer — logo e rodapé do template 433
    rebuild_header(doc.paragraphs[0])

    for i, topic in enumerate(TOPICS):
        base = 1 + i * 10
        set_plain(doc.paragraphs[base], topic["title"], bold=True, size_pt=17)
        for j, bullet in enumerate(topic["bullets"][:3]):
            set_plain(doc.paragraphs[base + 1 + j], f"🔹 {bullet}", size_pt=15)
        set_plain(doc.paragraphs[base + 4], "✔️ Benefícios:", size_pt=15)
        for j, benefit in enumerate(topic["benefits"][:4]):
            set_plain(doc.paragraphs[base + 5 + j], benefit, size_pt=15)

    numbered = [p for p in doc.paragraphs if re.match(r"^\d+\.\s+", (p.text or "").strip())]
    assert len(numbered) >= 7, len(numbered)

    labels = json.loads((SHOTS / "labels.json").read_text())
    img_idx = None
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip().startswith("Imagens"):
            img_idx = i
            break
    if img_idx is None:
        raise RuntimeError("Parágrafo Imagens não encontrado no template")

    set_plain(doc.paragraphs[img_idx], "Imagens", bold=True, size_pt=17)
    add_page_break_before(doc.paragraphs[img_idx])

    body = doc.paragraphs[img_idx]._p.getparent()
    after = doc.paragraphs[img_idx]._p.getnext()
    while after is not None:
        nxt = after.getnext()
        if after.tag == qn("w:sectPr"):
            break
        body.remove(after)
        after = nxt

    assert len(labels) == 4
    for n, item in enumerate(labels[:4], start=1):
        img_path = SHOTS / item["file"]
        assert img_path.exists(), img_path
        label = item["label"]
        if label.startswith("Imagem "):
            # labels.json já vem com prefixo
            caption = label
        else:
            caption = f"Imagem {n} — {label}"

        p_label = doc.add_paragraph()
        p_label.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_run(p_label, caption, bold=True, size_pt=12)

        p_pic = doc.add_paragraph()
        p_pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_pic.add_run().add_picture(str(img_path), width=picture_width_cm(img_path))

    doc.save(out_docx)
    clean_docx_zip(out_docx)

    doc2 = Document(out_docx)
    assert_center_title_and_images(doc2)
    assert_header_footer_branding(doc2)
    n = count_body_drawings(out_docx)
    assert n == 4, n

    text = "\n".join(p.text for p in doc2.paragraphs)
    assert "localhost" not in text.lower()
    assert LEGAL in text
    print(f"Generated {out_docx} drawings={n} header/footer OK")
    return out_docx


def main():
    out = build_checklist()
    token = get_token()
    year_id = drive_ensure_folder(PROJECT_FOLDER_ID, "2026", token)
    set_id = drive_ensure_folder(year_id, "Setembro", token)
    day_id = drive_ensure_folder(set_id, "01", token)
    drive_trash(day_id, [f"Checklist {PROJECT}", f"Checklist {PROJECT}.docx"], token)
    up = drive_upload(day_id, out, f"Checklist {PROJECT}", token)
    print("Uploaded Setembro/01", up.get("id"))
    print(f"https://drive.google.com/file/d/{up['id']}/view")


if __name__ == "__main__":
    main()
